from pathlib import Path
from zipfile import ZipFile
import re
from openpyxl import load_workbook

ROOT = Path(r"C:\Users\sammt\OneDrive\Documentos\ClimaPrevV3\climaprev-atualiza-ao-geral")
EXPECTED = {
    "importacao_apoiadores_Leva01.xlsx": 45,
    "importacao_apoiadores_LEVA04.xlsx": 31,
    "importacao_apoiadores_LEVA07.xlsx": 35,
}

def has_macro_entries(path: Path):
    with ZipFile(path) as zf:
        names = [n.lower() for n in zf.namelist()]
    needles = ("vbaproject", "customui", "activex", "ctrlprops", "vba", "macros")
    return [n for n in names if any(k in n for k in needles)]

def main():
    for name, expected_rows in EXPECTED.items():
        path = ROOT / name
        wb = load_workbook(path, data_only=False, read_only=False)
        ws = wb[wb.sheetnames[0]]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        macros = has_macro_entries(path)
        formulas = []
        questions = []
        blank_names = []
        invalid_sex = []
        title_bad_len = []
        cep_bad = []
        duplicate_names = []
        seen_names = {}
        for rnum, row in enumerate(rows, start=2):
            values = ["" if v is None else str(v) for v in row]
            for c, v in enumerate(row, start=1):
                if isinstance(v, str) and v.startswith("="):
                    formulas.append((rnum, c, v))
                if isinstance(v, str) and "?" in v:
                    questions.append((rnum, c, v))
            name_val = (row[0] or "").strip() if isinstance(row[0], str) else row[0]
            if not name_val:
                blank_names.append(rnum)
            else:
                seen_names.setdefault(name_val, []).append(rnum)
            if row[2] not in (None, "", "MASCULINO", "FEMININO", "OUTROS"):
                invalid_sex.append((rnum, row[2]))
            title = "" if row[10] is None else str(row[10])
            if title and not re.fullmatch(r"\d+", title):
                title_bad_len.append((rnum, title, "não numérico"))
            elif title and len(title) != 12:
                title_bad_len.append((rnum, title, f"{len(title)} dígitos"))
            cep = "" if row[21] is None else str(row[21])
            if cep and (not re.fullmatch(r"\d{8}", cep)):
                cep_bad.append((rnum, cep))
        duplicate_names = [(n, rs) for n, rs in seen_names.items() if len(rs) > 1]
        blank_by_col = {}
        for c, h in enumerate(headers, start=1):
            blank_by_col[str(h)] = sum(1 for row in rows if row[c-1] in (None, ""))
        print(f"[{name}]")
        print(f"  sheet={ws.title!r} rows={len(rows)} expected={expected_rows} cols={ws.max_column} expected_cols=32")
        print(f"  headers_32={len(headers)==32} row_count_ok={len(rows)==expected_rows} macros={macros or 'none'} formulas={len(formulas)}")
        print(f"  blank_names={blank_names or 'none'} invalid_sex={invalid_sex or 'none'} question_mark_values={questions or 'none'}")
        print(f"  title_non_12_digit={title_bad_len or 'none'}")
        print(f"  cep_invalid_format={cep_bad or 'none'} duplicate_names={duplicate_names or 'none'}")
        print("  blanks_by_column=" + "; ".join(f"{h}:{n}" for h, n in blank_by_col.items() if n))

if __name__ == "__main__":
    main()
