from copy import copy
from pathlib import Path
import shutil
import zipfile
import openpyxl

INPUT = Path(r"C:\Users\sammt\OneDrive\Área de Trabalho\modelo_importacao_apoiadores.xlsx")
OUTPUT = Path(r"C:\Users\sammt\OneDrive\Documentos\ClimaPrevV3\climaprev-atualiza-ao-geral\importacao_apoiadores_LEVA07.xlsx")

def rec(nome, sexo, titulo="", secao="", zona="", endereco="", n_casa="", cep="", bairro="", telefone="", complemento=""):
    return {"NOME": nome, "SEXO": sexo, "N TÍTULO": titulo, "SEÇÃO": secao, "ZONA": zona, "ENDEREÇO": endereco,
            "N CASA": n_casa, "CEP": cep, "COMPLEMENTO": complemento, "BAIRRO": bairro, "TELEFONE": telefone}

rows = [
    rec("PAULISSON PINHEIRO DE ARAUJO", "MASCULINO", "086980661376", "0567", "076", "PASS. OLIVEIRA BELO", "50", "66060445", "FATIMA", "91981272944"),
    rec("WANESSA DO SOCORRO BRITO", "FEMININO", "039368391347", "0098", "095", "PASSAGEM CIA ROTA", "27", "66085100", "PEDREIRA", "984381295"),
    rec("MARIA LUIZA RAMOS CARDOSO", "FEMININO", "082571691384", "0304", "095", "PASSAGEM CIA ROTA", "27", "66085100", "PEDREIRA", "984792848"),
    rec("MARILENA DAMASCENO DA CASTRO", "FEMININO", "073365061364", "0432", "095", "AV. VISCONDE DE INHAUMA", "854", "66087640", "PEDREIRA", "986297237"),
    rec("JOAO MACHADO DE ARAUJO FILHO", "MASCULINO", "070905451368", "0177", "095", "PASSAGEM OLIVEIRA BELO", "501", "66060445", "FATIMA", "982790052"),

    rec("JUAN CARLOS MESSIAS MENEZES ESTEVES", "MASCULINO", "038573213392", "0272", "095", "RUA DIOGO NOIA", "1632", "66060140", "FATIMA", "981723999"),
    rec("RUI CARLOS MENEZES ESTEVES", "MASCULINO", "038496613350", "0260", "096", "RUA DIOGO NOIA", "1632", "66060140", "FATIMA", "987493782"),
    rec("HELENA HELENA R. DA ANUNCIACAO", "FEMININO", "012274871392", "0047", "096", "TRAV. 3 DE MAIO", "247", "66060600", "FATIMA", "983660005"),
    rec("HELENA CRISTINA RIBEIRO LIMA", "FEMININO", "060841513309", "0278", "076", "TRAV. 3 DE MAIO", "247", "66060600", "FATIMA", "981238294"),
    rec("PAULO DAVIDES P. NUNES", "MASCULINO", "057847401317", "0277", "076", "TRAV. 3 DE MAIO", "247", "66060600", "FATIMA", "986112445"),

    rec("MARIA DO SOCORRO CORREA MARIA", "FEMININO", "013681016376", "0161", "095", "TRAV. 3 DE MAIO", "", "66060600", "FATIMA", "984023106"),
    rec("EDNA DE FATIMA RIBEIRO DE ALMEIDA", "FEMININO", "012272461392", "0625", "095", "TRAV. 3 DE MAIO", "", "66060600", "FATIMA", "99874907"),
    rec("M. DE NAZARE LIMA SILVA", "FEMININO", "043664971309", "0570", "076", "TRAV. 3 DE MAIO", "", "66060600", "FATIMA", "98064110"),
    rec("HENRIQUE D. DE ARAGUACAS", "MASCULINO", "", "0047", "096", "TRAV. 3 DE MAIO", "247", "66060600", "FATIMA", "983680054"),
    rec("REGINALDO FERREIRA DO MARACUJA", "MASCULINO", "03782098376", "0009", "096", "TRAV. 3 DE MAIO", "247", "66060600", "FATIMA", "98389849"),

    rec("THAIS CRISTINA FRAZAO CAVALCANTE", "FEMININO", "068327341309", "0216", "095", "PASS. OLIVEIRA BELO", "501", "66060445", "FATIMA", "91981771087"),
    rec("BRUNO NATAN CORDEIRO COSTA", "MASCULINO", "066896201376", "0235", "095", "PASS. OLIVEIRA BELO", "501", "66060445", "FATIMA", "91982892240"),
    rec("GLEIDE ARAUJO CAVALCANTE", "FEMININO", "", "0269", "076", "PASS. OLIVEIRA BELO", "501", "66060445", "FATIMA", "9199829907"),
    rec("GLEICE FRANCA CAVALCANTE", "FEMININO", "036727971309", "0281", "076", "PASS. OLIVEIRA BELO", "501", "66060445", "FATIMA", "9193291062"),
    rec("HAMILTON SOUTO DOS REIS", "MASCULINO", "012240362325", "0220", "095", "PASS. OLIVEIRA BELO", "501", "66060445", "FATIMA", "91981301346"),

    rec("SIMONE LIMA COSTA", "FEMININO", "063107303392", "0083", "028", "PASSAGEM GABRIEL", "467", "66113200", "PEDREIRA - TELEGRAFO", "91998352493"),
    rec("GUILHERME SANTOS DE AZEVEDO", "MASCULINO", "011116513384", "0104", "003", "PASS. OLIVEIRA BELO", "502", "66060445", "FATIMA", "91980261087"),
    rec("KARLA RAMOS DA SILVA", "FEMININO", "061043413350", "0225", "095", "PASS. CARLOS FLORES", "79", "66113420", "TELEGRAFO", "9198175518"),
    rec("DALMA COSTA DOS SANTOS", "FEMININO", "068339091376", "0216", "095", "PASS. CARLOS FLORES", "", "66113420", "TELEGRAFO", "91980424145"),
    rec("LUCINETE SANTOS DE AZEVEDO", "FEMININO", "011124251333", "0274", "076", "PASS. OLIVEIRA BELO", "501", "66060445", "FATIMA", "91982981047"),

    rec("MARIA DO SOCORRO MOREIRA LUIZ", "FEMININO", "03370541384", "0277", "076", "PASS. OLIVEIRA BELO - MARQUES E DIOGO MOIA", "", "66060445", "FATIMA", "9198175655"),
    rec("JESSICA HELENA CRUZ DOS SANTOS", "FEMININO", "059293571384", "0225", "095", "PASS. OLIVEIRA BELO", "500", "66060445", "FATIMA", "91981546775"),
    rec("ANDERSON EDMILIO FREITAS DO NASCIMENTO", "MASCULINO", "004869192526", "0225", "095", "PASS. OLIVEIRA BELO", "", "66060445", "FATIMA", "91986431892"),
    rec("ADRIANA CRISTINA DOS SANTOS", "FEMININO", "045479173392", "0061", "095", "RUA DO ACAMPAMENTO", "81", "66083030", "PEDREIRA", "91981956455"),
    rec("WALBER DOS SANTOS RODRIGUES", "MASCULINO", "050158701384", "0279", "076", "RUA RIBEIRO GUAMA", "22", "66060445", "FATIMA", "91981546775"),

    rec("DANIELE CRISTINA PINHEIRO RIBEIRO", "FEMININO", "045010941333", "0271", "076", "PASS. OLIVEIRA BELO", "", "66060445", "FATIMA", "9198219562"),
    rec("HOWERTON AUGUSTO CRUZ DOS SANTOS", "MASCULINO", "059290313376", "0230", "095", "PASS. OLIVEIRA BELO", "501", "66060445", "FATIMA", "91983017782"),
    rec("RAIMUNDA ANACILIA AMERICO COSTA", "FEMININO", "010958013376", "0222", "095", "PASS. OLIVEIRA BELO", "498", "66060445", "FATIMA", "91982156723"),
    rec("LEYCLE MARY SOARES JARDIM BARBOSA", "FEMININO", "033362771309", "0135", "095", "PASS. OLIVEIRA BELO", "504-A", "66060445", "FATIMA", "91985443534"),
    rec("ELIZE REIS DE MORAES", "FEMININO", "", "0061", "095", "RUA ROTA", "224", "66083543", "PEDREIRA", "91981546775"),
]

assert len(rows) == 35, len(rows)
with zipfile.ZipFile(INPUT) as zf:
    bad = [n for n in zf.namelist() if any(x in n.lower() for x in ("vbaproject", "macros", "activex", "ctrlprops"))]
if bad:
    raise RuntimeError(f"Modelo contém macros: {bad}")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
shutil.copy2(INPUT, OUTPUT)
wb = openpyxl.load_workbook(OUTPUT, keep_vba=False)
ws = wb[wb.sheetnames[0]]
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
styles = []
for c in range(1, ws.max_column + 1):
    cell = ws.cell(2, c)
    styles.append((copy(cell._style), cell.number_format, copy(cell.font), copy(cell.fill), copy(cell.border), copy(cell.alignment), copy(cell.protection)))
row_height = ws.row_dimensions[2].height
ws.delete_rows(2, ws.max_row - 1)
for data in rows:
    ws.append([data.get(str(headers[c - 1]), "") if headers[c - 1] is not None else "" for c in range(1, ws.max_column + 1)])
idx = {str(v): i + 1 for i, v in enumerate(headers) if v is not None}
text_cols = {idx.get(x) for x in ("N TÍTULO", "SEÇÃO", "ZONA", "N CASA", "CEP")}
for r in range(2, ws.max_row + 1):
    if row_height is not None: ws.row_dimensions[r].height = row_height
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(r, c)
        st, nf, font, fill, border, alignment, protection = styles[c - 1]
        cell._style = copy(st)
        cell.number_format = "@" if c in text_cols else nf
        cell.font, cell.fill, cell.border = copy(font), copy(fill), copy(border)
        cell.alignment, cell.protection = copy(alignment), copy(protection)
wb.save(OUTPUT)
with zipfile.ZipFile(OUTPUT) as zf:
    final_bad = [n for n in zf.namelist() if any(x in n.lower() for x in ("vbaproject", "macros", "activex", "ctrlprops"))]
if final_bad: raise RuntimeError(f"Arquivo final contém macros: {final_bad}")
check = openpyxl.load_workbook(OUTPUT, read_only=True)
assert check.active.max_row == 36 and check.active.max_column == 32
print(f"created={OUTPUT}; rows=35; columns=32; macros=[]")
